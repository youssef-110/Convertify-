"""Runs every conversion in the map on a sample file and checks the result opens."""
import json
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import converters  # noqa: E402
from converters.tools import ConversionError, find, run, status  # noqa: E402
from tests.make_samples import make  # noqa: E402

TOOLS = status()


@pytest.fixture(scope="session")
def samples():
    return make()


def check(path):
    """Open the output with the library for its type. Fails the test if it can't."""
    assert path.exists(), f"{path.name} was not created"
    assert path.stat().st_size > 0, f"{path.name} is empty"
    ext = path.suffix.lower().lstrip(".")
    if ext == "zip":
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            assert names
            z.extractall(path.parent / "_unzipped")
        for n in names:
            check(path.parent / "_unzipped" / n)
    elif ext == "pdf":
        import pymupdf as fitz
        doc = fitz.open(path)
        assert doc.needs_pass or doc.page_count > 0
    elif ext == "docx":
        import docx
        docx.Document(path)
    elif ext == "pptx":
        from pptx import Presentation
        assert len(Presentation(path).slides) > 0
    elif ext == "xlsx":
        import openpyxl
        assert openpyxl.load_workbook(path).sheetnames
    elif ext in ("odt", "ods", "odp"):
        with zipfile.ZipFile(path) as z:
            assert z.read("mimetype").startswith(b"application/vnd.oasis.opendocument")
    elif ext in converters.IMAGES:
        from PIL import Image
        from pillow_heif import register_heif_opener
        register_heif_opener()
        with Image.open(path) as im:
            im.load()
            assert im.width > 0
    elif ext in ("csv", "tsv"):
        import pandas as pd
        assert len(pd.read_csv(path, sep="\t" if ext == "tsv" else ",").columns) > 0
    elif ext == "epub":
        with zipfile.ZipFile(path) as z:
            assert z.read("mimetype") == b"application/epub+zip"
    elif ext == "json":
        assert json.loads(path.read_text(encoding="utf-8"))
    elif ext in ("txt", "html", "md"):
        assert path.read_text(encoding="utf-8").strip()
    elif ext in converters.VIDEO + converters.AUDIO:
        run([find("ffmpeg"), "-v", "error", "-i", path, "-f", "null", "-"], 60)  # decodes the whole file
    else:
        pytest.fail(f"No check for .{ext}")


CASES = [(ext, out) for ext, outs in converters.MAP.items() for out in outs if out != "pdf-unlock"]
OPTS = {"pdf-protect": {"password": "s3cret"}, "pdf-watermark": {"text": "DRAFT"}}  # actions that need input


@pytest.mark.parametrize("ext,target", CASES, ids=[f"{e}->{t}" for e, t in CASES])
def test_conversion(ext, target, samples, tmp_path):
    missing = [TOOLS[t]["name"] for t in converters.MAP[ext][target].needs if not TOOLS[t]["found"]]
    if missing:
        pytest.skip(f"needs {', '.join(missing)}")
    src = samples / f"sample.{ext}"
    if not src.exists():
        pytest.skip(f"no sample.{ext} (a program needed to create it is missing)")
    outs = converters.run(target, [src], tmp_path, {"_timeout": 300, **OPTS.get(target, {})})
    assert outs
    for o in outs:
        check(Path(o))


def test_ocr_reads_text(samples, tmp_path):
    if not TOOLS["tesseract"]["found"]:
        pytest.skip("needs Tesseract OCR")
    out = converters.run("ocr-txt", [samples / "sample.png"], tmp_path, {"_timeout": 120})[0]
    assert "12345" in out.read_text(encoding="utf-8")


def test_merge_pdfs(samples, tmp_path):
    out = converters.run("merge-pdf", [samples / "sample.pdf"] * 3, tmp_path, {})[0]
    import pymupdf as fitz
    assert fitz.open(out).page_count == 6


def test_images_to_one_pdf(samples, tmp_path):
    out = converters.run("images-pdf", [samples / "sample.png", samples / "sample.jpg", samples / "sample.heic"],
                         tmp_path, {})[0]
    import pymupdf as fitz
    assert fitz.open(out).page_count == 3


def test_split_with_page_range(samples, tmp_path):
    outs = converters.run("pdf-split", [samples / "sample.pdf"], tmp_path, {"pages": "2"})
    assert [o.name for o in outs] == ["sample_page2.pdf"]


def test_bad_page_range(samples, tmp_path):
    with pytest.raises(ConversionError, match="outside"):
        converters.run("txt", [samples / "sample.pdf"], tmp_path, {"pages": "5-9"})


def test_damaged_file_gives_clear_error(tmp_path):
    bad = tmp_path / "broken.png"
    bad.write_bytes(b"not an image")
    with pytest.raises(ConversionError, match="could not be opened"):
        converters.run("jpg", [bad], tmp_path, {})


def test_lock_then_unlock(samples, tmp_path):
    import pymupdf as fitz
    locked = converters.run("pdf-protect", [samples / "sample.pdf"], tmp_path, {"password": "s3cret"})[0]
    assert fitz.open(locked).needs_pass
    with pytest.raises(ConversionError, match="wrong"):
        converters.run("pdf-unlock", [locked], tmp_path, {"password": "nope"})
    with pytest.raises(ConversionError, match="no password"):
        converters.run("pdf-unlock", [samples / "sample.pdf"], tmp_path, {"password": "x"})
    unlocked = converters.run("pdf-unlock", [locked], tmp_path, {"password": "s3cret"})[0]
    assert fitz.open(unlocked).page_count == 2 and not fitz.open(unlocked).needs_pass


def test_protect_needs_password(samples, tmp_path):
    with pytest.raises(ConversionError, match="password"):
        converters.run("pdf-protect", [samples / "sample.pdf"], tmp_path, {})


def test_rotate_only_chosen_pages(samples, tmp_path):
    import pymupdf as fitz
    out = converters.run("pdf-rotate", [samples / "sample.pdf"], tmp_path, {"angle": "180", "pages": "2"})[0]
    doc = fitz.open(out)
    assert [p.rotation for p in doc] == [0, 180]


def test_compress_makes_pdf_smaller(samples, tmp_path):
    out = converters.run("pdf-compress", [samples / "sample.pdf"], tmp_path, {})[0]
    assert out.stat().st_size <= (samples / "sample.pdf").stat().st_size


def test_images_to_gif(samples, tmp_path):
    from PIL import Image
    out = converters.run("images-gif", [samples / "sample.png", samples / "sample.jpg"], tmp_path, {})[0]
    assert Image.open(out).n_frames == 2


def test_zip_and_unzip(samples, tmp_path):
    packed = converters.run("zip", [samples / "sample.txt", samples / "sample.txt"], tmp_path, {})[0]
    with zipfile.ZipFile(packed) as z:
        assert z.namelist() == ["sample.txt", "sample (2).txt"]
    (tmp_path / "x").mkdir()
    outs = converters.run("unzip", [samples / "sample.zip"], tmp_path / "x", {})
    assert sorted(o.name for o in outs) == ["data.csv", "readme.txt"]


def test_trim_bad_time(samples, tmp_path):
    if not TOOLS["ffmpeg"]["found"]:
        pytest.skip("needs FFmpeg")
    with pytest.raises(ConversionError, match="isn't a time"):
        converters.run("mp3", [samples / "sample.mp4"], tmp_path, {"start": "abc", "_timeout": 60})
